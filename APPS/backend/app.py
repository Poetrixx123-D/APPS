from flask import Flask, render_template, request, redirect, url_for, session, Response
import sqlite3
import os
import joblib
from datetime import date, datetime, timedelta
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
# Load the trained ML model (trained by train_model.py)
model = joblib.load("stress_model.pkl")
app.secret_key = "apps_dev_secret_key"

DB_PATH = os.path.join(os.path.dirname(__file__), "apps.db")


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS Users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            unit TEXT,
            password TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'personnel'
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS DailyData (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            date TEXT NOT NULL,
            sleep_hours REAL,
            duty_hours REAL,
            heart_rate INTEGER,
            mood_score INTEGER,
            stress_level TEXT,
            recommendation TEXT
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS OfficerCodes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT UNIQUE NOT NULL,
            is_used INTEGER NOT NULL DEFAULT 0
        )
    """)

    # Pre-load some unique officer codes (only runs if table is empty)
    existing = c.execute("SELECT COUNT(*) FROM OfficerCodes").fetchone()[0]
    if existing == 0:
        starter_codes = [
            "OFC-7X2K9",
            "OFC-4M8P1",
            "OFC-9R3T6",
            "OFC-2L5Q8",
            "OFC-6N1W4",
        ]
        for code in starter_codes:
            c.execute("INSERT INTO OfficerCodes (code, is_used) VALUES (?, 0)", (code,))

    c.execute("""
        CREATE TABLE IF NOT EXISTS StressRequests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            message TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Pending',
            stage TEXT NOT NULL DEFAULT 'New',
            officer_id INTEGER,
            remedy TEXT,
            created_date TEXT NOT NULL,
            resolved_date TEXT,
            stage_updated_at TEXT
        )
    """)

    # Migrate older databases that already had StressRequests without the
    # pipeline columns (stage tracking was added after the table existed).
    request_cols = [row["name"] for row in c.execute("PRAGMA table_info(StressRequests)").fetchall()]
    if "stage" not in request_cols:
        c.execute("ALTER TABLE StressRequests ADD COLUMN stage TEXT NOT NULL DEFAULT 'New'")
    if "stage_updated_at" not in request_cols:
        c.execute("ALTER TABLE StressRequests ADD COLUMN stage_updated_at TEXT")

    # Backfill: anything already marked Resolved under the old status-only
    # model starts at the "Resolved" stage instead of the pipeline's start.
    c.execute("UPDATE StressRequests SET stage = 'Resolved' WHERE status = 'Resolved' AND stage = 'New'")

    c.execute("""
        CREATE TABLE IF NOT EXISTS CounselingNotes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            officer_id INTEGER NOT NULL,
            date TEXT NOT NULL,
            notes TEXT,
            action_taken TEXT
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS AdminCodes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT UNIQUE NOT NULL,
            is_used INTEGER NOT NULL DEFAULT 0
        )
    """)

    existing_admin = c.execute("SELECT COUNT(*) FROM AdminCodes").fetchone()[0]
    if existing_admin == 0:
        admin_codes = [
            "ADM-1A9F3",
            "ADM-5B2H7",
            "ADM-8C4J0",
        ]
        for code in admin_codes:
            c.execute("INSERT INTO AdminCodes (code, is_used) VALUES (?, 0)", (code,))

    conn.commit()
    conn.close()

# ---------- Stress request pipeline (CRM-style staged workflow) ----------
# A request must move through these stages IN ORDER - each step only
# unlocks once the one before it is done. No skipping ahead.
STAGE_FLOW = ["New", "Acknowledged", "In Progress", "Resolved", "Closed"]

STAGE_INFO = {
    "New":          {"label": "New",              "next_label": "Acknowledge Request"},
    "Acknowledged": {"label": "Acknowledged",      "next_label": "Start Working On This"},
    "In Progress":  {"label": "In Progress",       "next_label": "Mark Resolved"},
    "Resolved":     {"label": "Resolved",          "next_label": "Close Case"},
    "Closed":       {"label": "Closed",            "next_label": None},
}


def get_next_stage(current_stage):
    """Returns the single next stage in the pipeline, or None if already at the end."""
    if current_stage not in STAGE_FLOW:
        return STAGE_FLOW[0]
    idx = STAGE_FLOW.index(current_stage)
    if idx + 1 < len(STAGE_FLOW):
        return STAGE_FLOW[idx + 1]
    return None


def calculate_stress(sleep_hours, duty_hours, heart_rate, mood_score):
    """
    Version 2: Uses the trained Random Forest ML model instead of rules.
    """
    prediction = model.predict([[sleep_hours, duty_hours, heart_rate, mood_score]])
    level = prediction[0]  # "Low", "Medium", or "High"

    recommendations = {
        "Low": "You're within a healthy range today. Keep up the balance of rest and duty.",
        "Medium": "Some strain showing up today. Try to get extra rest tonight if possible.",
        "High": "Your readings suggest significant strain. Please reach out to your welfare officer or medical staff."
    }
    recommendation = recommendations[level]

    return level, recommendation


@app.route("/")
def home():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        name = request.form["name"]
        unit = request.form["unit"]
        password = request.form["password"]

        conn = get_db()
        conn.execute(
            "INSERT INTO Users (name, unit, password, role) VALUES (?, ?, ?, ?)",
            (name, unit, password, "personnel")
        )
        conn.commit()
        conn.close()
        return redirect(url_for("login"))
    return render_template("signup.html")

@app.route("/officer_signup", methods=["GET", "POST"])
def officer_signup():
    error = None
    if request.method == "POST":
        name = request.form["name"]
        unit = request.form["unit"]
        password = request.form["password"]
        entered_code = request.form["access_code"].strip()

        conn = get_db()
        code_row = conn.execute(
            "SELECT * FROM OfficerCodes WHERE code = ? AND is_used = 0",
            (entered_code,)
        ).fetchone()

        if code_row:
            conn.execute(
                "INSERT INTO Users (name, unit, password, role) VALUES (?, ?, ?, ?)",
                (name, unit, password, "officer")
            )
            conn.execute(
                "UPDATE OfficerCodes SET is_used = 1 WHERE code = ?",
                (entered_code,)
            )
            conn.commit()
            conn.close()
            return redirect(url_for("login"))
        else:
            conn.close()
            error = "Invalid or already-used access code. Please contact your administrator."

    return render_template("officer_signup.html", error=error)


@app.route("/admin_signup", methods=["GET", "POST"])
def admin_signup():
    error = None
    if request.method == "POST":
        name = request.form["name"]
        unit = request.form["unit"]
        password = request.form["password"]
        entered_code = request.form["access_code"].strip()

        conn = get_db()
        code_row = conn.execute(
            "SELECT * FROM AdminCodes WHERE code = ? AND is_used = 0",
            (entered_code,)
        ).fetchone()

        if code_row:
            conn.execute(
                "INSERT INTO Users (name, unit, password, role) VALUES (?, ?, ?, ?)",
                (name, unit, password, "admin")
            )
            conn.execute(
                "UPDATE AdminCodes SET is_used = 1 WHERE code = ?",
                (entered_code,)
            )
            conn.commit()
            conn.close()
            return redirect(url_for("login"))
        else:
            conn.close()
            error = "Invalid or already-used access code. Please contact your administrator."

    return render_template("admin_signup.html", error=error)

@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        name = request.form["name"]
        password = request.form["password"]

        conn = get_db()
        user = conn.execute(
            "SELECT * FROM Users WHERE name = ? AND password = ?",
            (name, password)
        ).fetchone()
        conn.close()

        if user:
            session["user_id"] = user["id"]
            session["user_name"] = user["name"]
            session["role"] = user["role"]

            if user["role"] == "officer":
                return redirect(url_for("officer_dashboard"))
            elif user["role"] == "admin":
                return redirect(url_for("admin_dashboard"))
            else:
                return redirect(url_for("dashboard"))
        else:
            error = "Invalid name or password. Please try again or sign up."

    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------- Compare with unit average (Personnel dashboard, feature 6) ----------
# key, label, suffix, higher_is_better, "about the same" tolerance, bar max
COMPARE_METRICS = [
    ("sleep", "Sleep",      "hrs", True,  0.5, 12),
    ("duty",  "Duty",       "hrs", False, 0.5, 24),
    ("hr",    "Heart rate", "bpm", False, 3,   150),
    ("mood",  "Mood",       "/5",  True,  0.5, 5),
]
COMPARE_DAYS = 7          # look-back window
COMPARE_MIN_MEMBERS = 3   # privacy: never show a unit average built from fewer people


def get_unit_comparison(conn, user_id, unit):
    """
    Compare a person's last-7-day averages with their unit's last-7-day average.
    Each person is averaged first, then the unit is the mean of those person-averages,
    so someone who checks in twice a day doesn't outweigh everyone else.
    Only aggregates leave this function - no other person's data is exposed.
    """
    if not unit:
        return {"status": "no_unit"}

    since = (date.today() - timedelta(days=COMPARE_DAYS - 1)).isoformat()
    rows = conn.execute("""
        SELECT DailyData.user_id AS uid,
               AVG(DailyData.sleep_hours) AS sleep,
               AVG(DailyData.duty_hours)  AS duty,
               AVG(DailyData.heart_rate)  AS hr,
               AVG(DailyData.mood_score)  AS mood
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE Users.role = 'personnel'
          AND Users.unit = ?
          AND DailyData.date >= ?
        GROUP BY DailyData.user_id
    """, (unit, since)).fetchall()

    mine = next((r for r in rows if r["uid"] == user_id), None)
    if mine is None:
        return {"status": "no_checkins", "unit": unit}
    if len(rows) < COMPARE_MIN_MEMBERS:
        return {"status": "not_enough", "unit": unit, "min_members": COMPARE_MIN_MEMBERS}

    metrics = []
    watch_count = 0
    for key, label, suffix, higher_better, tol, bar_max in COMPARE_METRICS:
        you = mine[key]
        unit_avg = sum(r[key] for r in rows) / len(rows)
        diff = you - unit_avg

        if abs(diff) <= tol:
            verdict = "same"
        elif (diff > 0) == higher_better:
            verdict = "good"
        else:
            verdict = "watch"
            watch_count += 1

        metrics.append({
            "label": label,
            "suffix": suffix,
            "you": round(you, 1),
            "unit": round(unit_avg, 1),
            "diff": round(diff, 1),
            "verdict": verdict,
            "you_pct": min(100, round(you / bar_max * 100)),
            "unit_pct": min(100, round(unit_avg / bar_max * 100)),
        })

    if watch_count == 0:
        summary = "You're in line with, or doing better than, your unit this week."
    elif watch_count == 1:
        summary = "One indicator is worth a look this week - see the amber row below."
    else:
        summary = f"{watch_count} indicators are worth a look this week - see the amber rows below."

    return {
        "status": "ok",
        "unit": unit,
        "members": len(rows),
        "days": COMPARE_DAYS,
        "metrics": metrics,
        "summary": summary,
        "watch_count": watch_count,
    }


@app.route("/dashboard")
def dashboard():
    if "user_id" not in session:
        return redirect(url_for("login"))

    conn = get_db()
    history = conn.execute(
        "SELECT * FROM DailyData WHERE user_id = ? ORDER BY id DESC",
        (session["user_id"],)
    ).fetchall()

    # Compare with unit average (personnel only)
    compare = None
    if session.get("role") == "personnel":
        me = conn.execute("SELECT unit FROM Users WHERE id = ?", (session["user_id"],)).fetchone()
        compare = get_unit_comparison(conn, session["user_id"], me["unit"] if me else None)
    conn.close()

    latest = history[0] if history else None

        # Quick tips based on the latest check-in
    tips = []
    if latest:
        sleep = latest["sleep_hours"]
        duty = latest["duty_hours"]
        hr = latest["heart_rate"]
        mood = latest["mood_score"]

        if mood <= 2:
            tips.append("You're feeling low today. Consider using 'Talk to Welfare Officer' to share what's on your mind.")
        elif mood == 3:
            tips.append("Mood is so-so. A short walk or a chat with a colleague can lift it.")

        if sleep < 5:
            tips.append("You slept under 5 hours. Try to get to bed 45-60 minutes earlier tonight.")
        elif sleep < 6.5:
            tips.append("Your sleep was a little short. Aim for 7+ hours tonight and keep screens away for 30 minutes before bed.")

        if duty > 12:
            tips.append("Very long duty day. Take a proper meal and a short rest break, and rest well before your next shift.")
        elif duty > 9:
            tips.append("Heavy duty day. Short breaks between tasks help keep fatigue down.")

        if hr > 90:
            tips.append("Your resting heart rate is high. Try 5 minutes of slow breathing (4 seconds in, 6 out) and drink some water.")
        elif hr > 80:
            tips.append("Resting heart rate is slightly raised. A short walk or some stretching can help.")

        if not tips:
            tips.append("All your indicators look healthy today. Keep up the balance of rest and duty.")

        tips = tips[:3]

    streak = 0
    if history:
        expected = date.today()
        dates = {row["date"] for row in history}
        while expected.isoformat() in dates:
            streak += 1
            expected -= timedelta(days=1)

    # Milestone check for celebration banner
    MILESTONES = [7, 14, 30, 60, 100]
    milestone = streak if streak in MILESTONES else None

    # Check if today's check-in is missing
    today_str = date.today().isoformat()
    checked_in_today = any(row["date"] == today_str for row in history)

    # Fetch personnel's own stress request status (staged pipeline)
    req_conn = get_db()
    pending_request = req_conn.execute(
        "SELECT * FROM StressRequests WHERE user_id = ? AND stage NOT IN ('Resolved', 'Closed') ORDER BY id DESC LIMIT 1",
        (session["user_id"],)
    ).fetchone()
    resolved_request = None
    if not pending_request:
        resolved_request = req_conn.execute(
            "SELECT * FROM StressRequests WHERE user_id = ? AND stage IN ('Resolved', 'Closed') ORDER BY id DESC LIMIT 1",
            (session["user_id"],)
        ).fetchone()
    req_conn.close()

    hour = datetime.now().hour
    if hour < 12:
        greeting = "Good morning"
    elif hour < 17:
        greeting = "Good afternoon"
    else:
        greeting = "Good evening"

    # Last 7 check-ins, oldest -> newest, for the personal trend chart
    recent = list(history[:7])
    recent.reverse()
    trend_labels = [row["date"] for row in recent]
    trend_sleep = [row["sleep_hours"] for row in recent]
    trend_duty = [row["duty_hours"] for row in recent]

    return render_template(
        "dashboard.html",
        user_name=session["user_name"],
        history=history,
        latest=latest,
        streak=streak,
        milestone=milestone,
        checked_in_today=checked_in_today,
        pending_request=pending_request,
        resolved_request=resolved_request,
        tips=tips,
        greeting=greeting,
        today=date.today().isoformat(),
        trend_labels=trend_labels,
        trend_sleep=trend_sleep,
        trend_duty=trend_duty,
        compare=compare
    )

@app.route("/results")
def results():
    if "user_id" not in session:
        return redirect(url_for("login"))

    conn = get_db()
    history = conn.execute(
        "SELECT * FROM DailyData WHERE user_id = ? ORDER BY id DESC",
        (session["user_id"],)
    ).fetchall()
    conn.close()

    return render_template(
        "results.html",
        user_name=session["user_name"],
        history=history
    )

@app.route("/officer_dashboard")
def officer_dashboard():
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    # Get every personnel's most recent check-in
    personnel = conn.execute("""
        SELECT Users.id, Users.name, Users.unit, DailyData.date, DailyData.sleep_hours,
               DailyData.duty_hours, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        ORDER BY DailyData.stress_level = 'High' DESC, DailyData.stress_level = 'Medium' DESC
    """).fetchall()
    conn.close()

    total = len(personnel)
    high_count = sum(1 for p in personnel if p["stress_level"] == "High")
    medium_count = sum(1 for p in personnel if p["stress_level"] == "Medium")
    low_count = sum(1 for p in personnel if p["stress_level"] == "Low")

    return render_template(
        "officer_dashboard.html",
        user_name=session["user_name"],
        personnel=personnel,
        total=total,
        high_count=high_count,
        medium_count=medium_count,
        low_count=low_count
    )

@app.route("/admin_dashboard")
def admin_dashboard():
    if "user_id" not in session or session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    # Latest check-in per personnel, joined with their unit
    rows = conn.execute("""
        SELECT Users.unit, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        AND Users.role = 'personnel'
    """).fetchall()
    conn.close()

    # Build unit-wise summary in Python
    unit_summary = {}
    for row in rows:
        unit = row["unit"] or "Unassigned"
        level = row["stress_level"]
        if unit not in unit_summary:
            unit_summary[unit] = {"total": 0, "High": 0, "Medium": 0, "Low": 0}
        unit_summary[unit]["total"] += 1
        unit_summary[unit][level] += 1

    total_personnel = len(rows)
    total_high = sum(1 for r in rows if r["stress_level"] == "High")
    total_medium = sum(1 for r in rows if r["stress_level"] == "Medium")
    total_low = sum(1 for r in rows if r["stress_level"] == "Low")

    return render_template(
        "admin_dashboard.html",
        admin_name=session["user_name"],
        unit_summary=unit_summary,
        total_personnel=total_personnel,
        total_high=total_high,
        total_medium=total_medium,
        total_low=total_low
    )


@app.route("/admin/export_report")
def export_report():
    if "user_id" not in session or session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    rows = conn.execute("""
        SELECT Users.unit, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        AND Users.role = 'personnel'
    """).fetchall()
    conn.close()

    unit_summary = {}
    for row in rows:
        unit = row["unit"] or "Unassigned"
        level = row["stress_level"]
        if unit not in unit_summary:
            unit_summary[unit] = {"total": 0, "High": 0, "Medium": 0, "Low": 0}
        unit_summary[unit]["total"] += 1
        unit_summary[unit][level] += 1

    total_personnel = len(rows)
    total_high = sum(1 for r in rows if r["stress_level"] == "High")
    total_medium = sum(1 for r in rows if r["stress_level"] == "Medium")
    total_low = sum(1 for r in rows if r["stress_level"] == "Low")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("APPS - Organization Stress Report", styles['Title']))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 16))

    summary_data = [
        ["Total Personnel", "High Risk", "Medium Risk", "Low Risk"],
        [str(total_personnel), str(total_high), str(total_medium), str(total_low)]
    ]
    summary_table = Table(summary_data, hAlign='LEFT')
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Unit-Wise Breakdown", styles['Heading2']))
    elements.append(Spacer(1, 8))

    table_data = [["Unit", "Total", "High", "Medium", "Low"]]
    for unit, counts in unit_summary.items():
        table_data.append([unit, counts["total"], counts["High"], counts["Medium"], counts["Low"]])

    unit_table = Table(table_data, repeatRows=1, hAlign='LEFT')
    unit_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(unit_table)

    doc.build(elements)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype='application/pdf',
        headers={'Content-Disposition': 'attachment; filename=admin_report.pdf'}
    )

@app.route("/admin/export_excel")
def export_excel():
    if "user_id" not in session or session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    rows = conn.execute("""
        SELECT Users.unit, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        AND Users.role = 'personnel'
    """).fetchall()
    conn.close()

    unit_summary = {}
    for row in rows:
        unit = row["unit"] or "Unassigned"
        level = row["stress_level"]
        if unit not in unit_summary:
            unit_summary[unit] = {"total": 0, "High": 0, "Medium": 0, "Low": 0}
        unit_summary[unit]["total"] += 1
        unit_summary[unit][level] += 1

    total_personnel = len(rows)
    total_high = sum(1 for r in rows if r["stress_level"] == "High")
    total_medium = sum(1 for r in rows if r["stress_level"] == "Medium")
    total_low = sum(1 for r in rows if r["stress_level"] == "Low")

    wb = Workbook()
    ws = wb.active
    ws.title = "Organization Summary"

    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws["A1"] = "APPS - Organization Stress Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws.append([])

    ws.append(["Total Personnel", "High Risk", "Medium Risk", "Low Risk"])
    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
    ws.append([total_personnel, total_high, total_medium, total_low])
    ws.append([])

    ws.append(["Unit", "Total", "High", "Medium", "Low"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
    for unit, counts in unit_summary.items():
        ws.append([unit, counts["total"], counts["High"], counts["Medium"], counts["Low"]])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=admin_report.xlsx'}
    )

@app.route("/officer/export_report")
def officer_export_report():
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    personnel = conn.execute("""
        SELECT Users.name, Users.unit, DailyData.date, DailyData.sleep_hours,
               DailyData.duty_hours, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        ORDER BY DailyData.stress_level = 'High' DESC, DailyData.stress_level = 'Medium' DESC
    """).fetchall()
    conn.close()

    total = len(personnel)
    high_count = sum(1 for p in personnel if p["stress_level"] == "High")
    medium_count = sum(1 for p in personnel if p["stress_level"] == "Medium")
    low_count = sum(1 for p in personnel if p["stress_level"] == "Low")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("APPS - Personnel Stress Report", styles['Title']))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 16))

    summary_data = [
        ["Total Personnel", "High Risk", "Medium Risk", "Low Risk"],
        [str(total), str(high_count), str(medium_count), str(low_count)]
    ]
    summary_table = Table(summary_data, hAlign='LEFT')
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Personnel List", styles['Heading2']))
    elements.append(Spacer(1, 8))

    table_data = [["Name", "Unit", "Date", "Sleep Hrs", "Duty Hrs", "Stress Level"]]
    for p in personnel:
        table_data.append([
            p["name"], p["unit"] or "-", p["date"] or "-",
            p["sleep_hours"], p["duty_hours"], p["stress_level"]
        ])

    personnel_table = Table(table_data, repeatRows=1, hAlign='LEFT')
    personnel_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(personnel_table)

    doc.build(elements)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype='application/pdf',
        headers={'Content-Disposition': 'attachment; filename=officer_report.pdf'}
    )


@app.route("/officer/export_excel")
def officer_export_excel():
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    personnel = conn.execute("""
        SELECT Users.name, Users.unit, DailyData.date, DailyData.sleep_hours,
               DailyData.duty_hours, DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        ORDER BY DailyData.stress_level = 'High' DESC, DailyData.stress_level = 'Medium' DESC
    """).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel Report"

    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws["A1"] = "APPS - Personnel Stress Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws.append([])

    ws.append(["Name", "Unit", "Date", "Sleep Hrs", "Duty Hrs", "Stress Level"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font

    for p in personnel:
        ws.append([p["name"], p["unit"] or "-", p["date"] or "-", p["sleep_hours"], p["duty_hours"], p["stress_level"]])

    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=officer_report.xlsx'}
    )

@app.route("/admin/unit/<unit_name>")
def admin_unit_detail(unit_name):
    if "user_id" not in session or session.get("role") != "admin":
        return redirect(url_for("login"))

    conn = get_db()
    # Latest check-in per personnel in this unit
    personnel = conn.execute("""
        SELECT Users.id, Users.name, Users.unit, DailyData.date, DailyData.sleep_hours,
               DailyData.duty_hours, DailyData.heart_rate, DailyData.mood_score,
               DailyData.stress_level
        FROM DailyData
        JOIN Users ON DailyData.user_id = Users.id
        WHERE DailyData.id IN (
            SELECT MAX(id) FROM DailyData GROUP BY user_id
        )
        AND Users.role = 'personnel'
        AND Users.unit = ?
        ORDER BY DailyData.stress_level = 'High' DESC, DailyData.stress_level = 'Medium' DESC
    """, (unit_name,)).fetchall()
    conn.close()

    return render_template(
        "admin_unit_detail.html",
        admin_name=session["user_name"],
        unit_name=unit_name,
        personnel=personnel
    )

@app.route("/officer/personnel/<int:user_id>")
def personnel_detail(user_id):
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    person = conn.execute("SELECT * FROM Users WHERE id = ?", (user_id,)).fetchone()
    history = conn.execute(
        "SELECT * FROM DailyData WHERE user_id = ? ORDER BY id DESC",
        (user_id,)
    ).fetchall()
    notes = conn.execute(
        """SELECT CounselingNotes.*, Users.name AS officer_name
           FROM CounselingNotes
           JOIN Users ON CounselingNotes.officer_id = Users.id
           WHERE CounselingNotes.user_id = ?
           ORDER BY CounselingNotes.id DESC""",
        (user_id,)
    ).fetchall()
    conn.close()

    # Last 10 check-ins, oldest -> newest, for the trend chart
    recent = list(history[:10])
    recent.reverse()
    trend_labels = [row["date"] for row in recent]
    trend_sleep = [row["sleep_hours"] for row in recent]
    trend_duty = [row["duty_hours"] for row in recent]

    # Group check-in history: Today / This Week / Older
    from datetime import timedelta
    today_str = date.today().isoformat()
    week_ago_str = (date.today() - timedelta(days=7)).isoformat()

    history_today = []
    history_week = []
    history_older = []
    for entry in history:
        if entry["date"] == today_str:
            history_today.append(entry)
        elif entry["date"] >= week_ago_str:
            history_week.append(entry)
        else:
            history_older.append(entry)

    return render_template(
        "personnel_detail.html",
        officer_name=session["user_name"],
        person=person,
        history=history,
        history_today=history_today,
        history_week=history_week,
        history_older=history_older,
        notes=notes,
        trend_labels=trend_labels,
        trend_sleep=trend_sleep,
        trend_duty=trend_duty
    )


@app.route("/add_note", methods=["POST"])
def add_note():
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    target_user_id = request.form["user_id"]
    note_text = request.form["notes"]
    action_taken = request.form["action_taken"]

    conn = get_db()
    conn.execute("""
        INSERT INTO CounselingNotes (user_id, officer_id, date, notes, action_taken)
        VALUES (?, ?, ?, ?, ?)
    """, (target_user_id, session["user_id"], date.today().isoformat(), note_text, action_taken))
    conn.commit()
    conn.close()

    return redirect(url_for("personnel_detail", user_id=target_user_id, saved=1))

@app.route("/submit_survey", methods=["POST"])
def submit_survey():
    if "user_id" not in session:
        return redirect(url_for("login"))

    sleep_hours = float(request.form["sleep_hours"])
    duty_hours = float(request.form["duty_hours"])
    heart_rate = int(request.form["heart_rate"])
    mood_score = int(request.form["mood_score"])

    level, recommendation = calculate_stress(sleep_hours, duty_hours, heart_rate, mood_score)

    conn = get_db()
    conn.execute("""
        INSERT INTO DailyData (user_id, date, sleep_hours, duty_hours, heart_rate, mood_score, stress_level, recommendation)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (session["user_id"], date.today().isoformat(), sleep_hours, duty_hours, heart_rate, mood_score, level, recommendation))
    conn.commit()
    conn.close()

    return redirect(url_for("dashboard"))


    # ---------- Route 1: Personnel submits a stress request ----------
@app.route("/submit_stress_request", methods=["POST"])
def submit_stress_request():
    if "user_id" not in session or session.get("role") != "personnel":
        return redirect(url_for("login"))

    conn = get_db()

    # Block a new request if one is already active anywhere in the pipeline
    # (anything short of Closed counts as "still being handled").
    existing = conn.execute(
        "SELECT id FROM StressRequests WHERE user_id = ? AND stage != 'Closed'",
        (session["user_id"],)
    ).fetchone()

    if not existing:
        message = request.form["message"]
        conn.execute(
            """INSERT INTO StressRequests (user_id, message, status, stage, created_date, stage_updated_at)
               VALUES (?, ?, 'Pending', 'New', ?, ?)""",
            (session["user_id"], message, date.today().isoformat(), datetime.now().isoformat())
        )
        conn.commit()

    conn.close()
    return redirect(url_for("dashboard"))


# ---------- Route 2: Officer views the request pipeline ----------
@app.route("/officer/requests")
def view_stress_requests():
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    # Everything still moving through the pipeline (not yet Closed)
    active_requests = conn.execute(
        """SELECT StressRequests.*, Users.name AS personnel_name
           FROM StressRequests
           JOIN Users ON StressRequests.user_id = Users.id
           WHERE StressRequests.stage != 'Closed'
           ORDER BY StressRequests.id DESC"""
    ).fetchall()
    conn.close()

    # Attach each request's next stage + button label for the template
    pending_requests = []
    for req in active_requests:
        row = dict(req)
        row["current_label"] = STAGE_INFO.get(row["stage"], STAGE_INFO["New"])["label"]
        row["next_stage"] = get_next_stage(row["stage"])
        row["next_label"] = STAGE_INFO.get(row["stage"], STAGE_INFO["New"])["next_label"]
        pending_requests.append(row)

    return render_template(
        "officer_requests.html",
        officer_name=session["user_name"],
        pending_requests=pending_requests,
        stage_flow=STAGE_FLOW
    )


# ---------- Route 3: Officer advances a request exactly one stage ----------
@app.route("/officer/requests/<int:request_id>/advance", methods=["POST"])
def advance_stress_request(request_id):
    if "user_id" not in session or session.get("role") != "officer":
        return redirect(url_for("login"))

    conn = get_db()
    req = conn.execute("SELECT * FROM StressRequests WHERE id = ?", (request_id,)).fetchone()
    if not req:
        conn.close()
        return redirect(url_for("view_stress_requests"))

    current_stage = req["stage"] or "New"
    expected_next = get_next_stage(current_stage)
    submitted_target = request.form.get("target_stage")

    # Guard rail: only allow moving to the exact next stage in the pipeline.
    # This is what enforces "finish this step before the next one starts" -
    # a stale form (or a tampered one) pointing at the wrong stage is ignored.
    if not expected_next or submitted_target != expected_next:
        conn.close()
        return redirect(url_for("view_stress_requests"))

    remedy = req["remedy"]
    resolved_date = req["resolved_date"]

    if expected_next == "Resolved":
        remedy = request.form.get("remedy", "").strip()
        if not remedy:
            # Can't complete this step without a remedy - stay put.
            conn.close()
            return redirect(url_for("view_stress_requests"))
        resolved_date = date.today().isoformat()

    status = "Resolved" if expected_next in ("Resolved", "Closed") else "Pending"

    conn.execute(
        """UPDATE StressRequests
           SET stage = ?, stage_updated_at = ?, officer_id = ?, remedy = ?,
               resolved_date = ?, status = ?
           WHERE id = ?""",
        (expected_next, datetime.now().isoformat(), session["user_id"],
         remedy, resolved_date, status, request_id)
    )
    conn.commit()
    conn.close()

    return redirect(url_for("view_stress_requests"))


if __name__ == "__main__":
    init_db()
    print("Server starting...")
    app.run(debug=True)